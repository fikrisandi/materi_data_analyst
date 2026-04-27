"""Convert dataset Excel ke SQLite database untuk latihan SQL.

Output: data/kopi_kita.db dengan 4 tabel:
- transaksi (250 baris)
- menu (8 baris)
- cabang (2 baris)
- pelanggan (50 baris)
"""

import sqlite3
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(__file__).parent.parent / "day-2-am-excel-foundation" / "data" / "penjualan-warung-2026.xlsx"
OUT = Path(__file__).parent / "data" / "kopi_kita.db"
OUT.parent.mkdir(parents=True, exist_ok=True)
if OUT.exists():
    OUT.unlink()

print(f"Reading from: {SRC}")
wb = load_workbook(SRC, read_only=True)
conn = sqlite3.connect(OUT)
cur = conn.cursor()

# ===== Schema =====
cur.executescript("""
CREATE TABLE cabang (
    id_cabang TEXT PRIMARY KEY,
    nama_cabang TEXT,
    kota TEXT,
    jam_buka TEXT,
    jam_tutup TEXT
);

CREATE TABLE menu (
    id_menu TEXT PRIMARY KEY,
    nama_menu TEXT,
    kategori TEXT,
    harga INTEGER,
    hpp INTEGER
);

CREATE TABLE pelanggan (
    id_pelanggan TEXT PRIMARY KEY,
    nama TEXT,
    no_hp TEXT,
    kota_asal TEXT,
    tanggal_join TEXT,
    membership TEXT
);

CREATE TABLE transaksi (
    id_transaksi TEXT PRIMARY KEY,
    tanggal TEXT,
    id_cabang TEXT,
    id_pelanggan TEXT,
    id_menu TEXT,
    qty INTEGER,
    harga_satuan INTEGER,
    total INTEGER,
    metode_bayar TEXT,
    FOREIGN KEY (id_cabang) REFERENCES cabang(id_cabang),
    FOREIGN KEY (id_pelanggan) REFERENCES pelanggan(id_pelanggan),
    FOREIGN KEY (id_menu) REFERENCES menu(id_menu)
);
""")

def insert_from_sheet(table_name, sheet_name, columns):
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    placeholders = ",".join(["?"] * len(columns))
    sql = f"INSERT INTO {table_name} ({','.join(columns)}) VALUES ({placeholders})"
    cleaned_rows = []
    for row in rows:
        if not row or row[0] is None: continue
        cleaned = []
        for v in row:
            if hasattr(v, "isoformat"):
                cleaned.append(v.isoformat())
            else:
                cleaned.append(v)
        cleaned_rows.append(tuple(cleaned[:len(columns)]))
    cur.executemany(sql, cleaned_rows)
    print(f"  {table_name}: inserted {len(cleaned_rows)} rows")

insert_from_sheet("cabang", "Cabang",
    ["id_cabang","nama_cabang","kota","jam_buka","jam_tutup"])
insert_from_sheet("menu", "Menu",
    ["id_menu","nama_menu","kategori","harga","hpp"])
insert_from_sheet("pelanggan", "Pelanggan",
    ["id_pelanggan","nama","no_hp","kota_asal","tanggal_join","membership"])
insert_from_sheet("transaksi", "Transaksi",
    ["id_transaksi","tanggal","id_cabang","id_pelanggan","id_menu","qty","harga_satuan","total","metode_bayar"])

conn.commit()

# verify
print("\nVerification:")
for tbl in ["cabang", "menu", "pelanggan", "transaksi"]:
    cur.execute(f"SELECT COUNT(*) FROM {tbl}")
    count = cur.fetchone()[0]
    print(f"  {tbl}: {count} rows")

conn.close()
print(f"\nSQLite database saved: {OUT}")
